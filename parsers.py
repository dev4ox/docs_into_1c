import os
import pandas as pd
from pathlib import Path
import re
import pdfplumber
import collections
import pytesseract
from pdf2image import convert_from_path
import settings


final_columns = ["Номенклатура", "Мощность, Вт", "Св. поток, Лм", "IP", "Габариты", "Длина, мм",
                 "Ширина, мм", "Высота, мм", "Рассеиватель", "Цвет. температура, К", "Вес, кг",
                 "Напряжение, В", "Температура эксплуатации", "Срок службы (работы) светильника",
                 "Тип КСС", "Род тока", "Гарантия", "Индекс цветопередачи (CRI, Ra)", "Цвет корпуса",
                 "Коэффициент пульсаций", "Коэффициент мощности (Pf)", "Класс взрывозащиты (Ex)",
                 "Класс пожароопасности", "Класс защиты от поражения электрическим током",
                 "Материал корпуса", "Тип", "Прочее"]


# Функция для извлечения текста из PDF
def extract_text_from_pdf(file_path: Path) -> str | None:
    if not file_path.exists():
        print(f"Файл не найден: {file_path}")
        return None
    text = ""
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        print("pdfplumber error:", e)
        return None

    # Если текст слишком короткий, попробуем OCR
    if len(text.strip()) < 100:
        try:
            pages = convert_from_path(str(file_path))
            for page in pages:
                text += pytesseract.image_to_string(page, lang='rus') + "\n"
        except Exception as e:
            print("OCR error:", e)
    return text.strip()


# Функция для извлечения текста из DOC/DOCX
def extract_text_from_docx(file_path: Path) -> str | None:
    if not file_path.exists():
        print(f"Файл не найден: {file_path}")
        return None
    try:
        from docx import Document
        doc = Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs])
    except Exception as e:
        print(f"Ошибка при обработке файла {file_path}: {e}")
        return None


class StructuredPdfParser:
    PRODUCT_NAMES = settings.product_names

    def __init__(self, file_path):
        self.file_path = Path(file_path)
        self.data = []

    def parse_pdf(self):
        print(f"Opening file: {self.file_path}")
        with pdfplumber.open(self.file_path) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                tables = page.extract_tables()
                if not tables:
                    continue

                for table_idx, table in enumerate(tables):
                    print(f"Processing table {table_idx + 1} on page {page_number}")
                    for row in table:
                        row_text = [str(cell).strip().replace("\n", " ") for cell in row if cell]
                        row_combined = " | ".join(row_text)

                        # Проверяем, содержится ли в строке название товара
                        if any(name.lower() in row_combined.lower() for name in self.PRODUCT_NAMES):
                            self.data.append(row_combined)

    def print_data(self) -> str | None:
        """Вывод полученных данных."""
        if not self.data:
            return "No data parsed!"
        else:
            return "\n".join(self.data)

    def process(self):
        if not self.file_path.exists():
            return f"File not found: {self.file_path}"

        print(f"Processing file: {self.file_path.name}")
        self.parse_pdf()



class UnifiedExcelParser:
    PRODUCT_NAMES = settings.product_names  # список названий товаров из настроек

    def __init__(self, file_path):
        self.file_path = Path(file_path)
        self.data = []

    def is_product_name(self, text):
        """Проверяет, начинается ли строка с одного из имен товаров."""
        return any(text.lower().startswith(name.lower()) for name in self.PRODUCT_NAMES)

    def detect_engine(self):
        """Определяет движок для pd.read_excel по расширению файла."""
        ext = self.file_path.suffix.lower()
        if ext in ['.xlsx', '.xlsm']:
            return 'openpyxl'
        elif ext == '.xls':
            return 'xlrd'
        else:
            return 'openpyxl'

    def parse_excel(self):
        """Главный метод, определяющий стратегию парсинга в зависимости от структуры файла."""
        if not self.file_path.exists():
            print(f"File not found: {self.file_path}")
            return
        engine = self.detect_engine()
        print(f"Opening file: {self.file_path} with engine {engine}")
        df = pd.read_excel(self.file_path, header=None, engine=engine)

        # Если в файле только один столбец, применяем логику, аналогичную варианту 2
        if df.shape[1] == 1:
            self.parse_single_column(df)
        else:
            # Поиск столбца, содержащего "наименование" в первых 15 строках
            name_column = None
            for col in df.columns:
                for row_idx in range(min(15, len(df))):
                    cell_value = str(df.iloc[row_idx, col]).lower() if pd.notna(df.iloc[row_idx, col]) else ""
                    if "наименование" in cell_value:
                        name_column = col
                        print(f"Found 'Наименование' column at index {col}")
                        break
                if name_column is not None:
                    break

            if name_column is None:
                # Если столбец не найден, пробуем обработку как одностолбцового файла
                print("Столбец 'Наименование' не найден, пробуем обработку как одностолбцовый файл.")
                self.parse_single_column(df)
            else:
                # Если рядом с найденным столбцом есть дополнительные столбцы – используем мультиколоночный парсинг
                if name_column + 1 < df.shape[1]:
                    # Если имеется ещё и столбец после characteristics (name_column+2), то считаем, что данные о характеристиках распределены на два столбца (как в варианте 3)
                    extra_char = (name_column + 2 < df.shape[1])
                    self.parse_multi_column(df, name_column, extra_char)
                else:
                    self.parse_single_column(df)

    def parse_single_column(self, df):
        """
        Обработка файла с одним столбцом (логика варианта 2):
        Каждая строка проверяется на то, является ли она названием нового товара.
        """
        current_product = None
        product_line = ""
        for index, row in df.iterrows():
            cell_value = str(row[0]).strip()
            if self.is_product_name(cell_value):
                if current_product:
                    self.data.append(product_line)
                current_product = cell_value
                product_line = current_product
            else:
                product_line += f" {cell_value}"  # Добавляем текст в одну строку
        if current_product:
            self.data.append(product_line)

    def parse_multi_column(self, df, name_column, extra_char):
        """
        Обработка файлов с несколькими столбцами (логика вариантов 1, 3, 4 и 5).
        Если extra_char==True – предполагается, что характеристики распределены на два столбца.
        """
        current_product = None
        product_line = ""
        for index, row in df.iterrows():
            if pd.isna(row[name_column]):
                continue
            cell_value = str(row[name_column]).strip().replace("\t", " ")
            char_value = ""
            if extra_char:
                if name_column + 1 < df.shape[1]:
                    char_value += str(row[name_column + 1]).strip().replace("\t", " ").replace("\n", ";")
                if name_column + 2 < df.shape[1]:
                    char_value += " " + str(row[name_column + 2]).strip().replace("\t", " ").replace("\n", ";")
            else:
                if name_column + 1 < df.shape[1]:
                    char_value = str(row[name_column + 1]).strip().replace("\t", " ").replace("\n", ";")

            if any(name.lower() in cell_value.lower() for name in self.PRODUCT_NAMES):
                if current_product:
                    self.data.append(product_line)
                current_product = cell_value
                product_line = current_product
                if char_value:
                    product_line += f" {char_value}"
            else:
                product_line += f" {cell_value}"  # Добавляем дополнительную информацию в строку
        if current_product:
            self.data.append(product_line)

    def print_data(self) -> str | None:
        """Вывод полученных данных."""
        if not self.data:
            return "No data parsed!"
        else:
            return "\n".join(self.data)

    def process(self):
        print()
        """Основной метод для обработки файла."""
        if not self.file_path.exists():
            return f"File not found: {self.file_path}"
        print(f"Processing file: {self.file_path.name}")
        self.parse_excel()



def append_df_to_excel(filename, df, sheet_name='Sheet1'):
    """
    Добавляет DataFrame в существующий Excel-файл.
    Если файл существует, новые строки дописываются в конец.
    """
    if os.path.exists(filename):
        with pd.ExcelWriter(filename, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
            startrow = writer.book[sheet_name].max_row
            df.to_excel(writer, sheet_name=sheet_name, index=False, header=False, startrow=startrow)
    else:
        df.to_excel(filename, index=False, sheet_name=sheet_name)