import pandas as pd
import docx
from openpyxl import load_workbook
import os
import re
from pathlib import Path


class StructuredDocxParser:
    def __init__(self, file_path):
        self.file_path = Path(file_path)
        self.data = []

    def parse_docx(self):
        print(f"Opening file: {self.file_path}")
        doc = docx.Document(self.file_path)

        all_data = []
        current_product = None
        product_data = {}

        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                row_combined = " | ".join(row_text)  # Разделяем ячейки для наглядности
                print(f"Row text: {row_combined}")  # Отладочный вывод

                # Проверяем, является ли строка номером нового товара (пример: 2.5)
                if re.match(r"^\d+\.\d+$", row_text[0]):
                    if current_product:
                        all_data.append(product_data)
                    current_product = row_combined
                    product_data = {"Номенклатура": current_product}
                # Проверяем, является ли строка характеристикой товара (пример: 2.5.1, 2.5.2)
                elif re.match(r"^\d+\.\d+\.\d+$", row_text[0]):
                    product_data[row_text[0]] = row_combined
                else:
                    product_data[f"Характеристика {len(product_data)}"] = row_combined

        if current_product:
            all_data.append(product_data)

        self.data = all_data

    def print_data(self):
        if not self.data:
            print("No data parsed!")
        for product in self.data:
            print(product)

    def process(self):
        if not self.file_path.exists():
            print(f"File not found: {self.file_path}")
            return

        print(f"Processing file: {self.file_path.name}")
        self.parse_docx()
        self.print_data()


# Использование
file_path = Path("..", "test_data", "input", "ТЗ для МГУ.docx")  # Конкретный файл DOCX
parser = StructuredDocxParser(file_path)
parser.process()