from pathlib import Path

from openpyxl import load_workbook
import pandas as pd
from spire.doc import Document
from spire.doc import FileFormat


def resize_column_in_intermediate_xlsx(path: Path) -> None:
    # 🔹 Загружаем созданный файл
    book = load_workbook(path)
    sheet = book.active

    # 🔹 Автоматическая ширина колонок
    for col in sheet.columns:
        max_length = 0
        col_letter = col[0].column_letter  # Получаем букву колонки (A, B, C...)

        for cell in col:
            try:
                max_length = max(max_length, len(str(cell.value)))  # Максимальная длина значения

            except:
                pass

        adjusted_width = max_length + 2  # Небольшой запас по ширине
        sheet.column_dimensions[col_letter].width = adjusted_width  # Устанавливаем ширину

    # 🔹 Сохраняем изменения
    book.save(path)


def create_intermediate_xlsx(path: Path) -> None:
    df = pd.DataFrame([], columns=["name", "value"])
    df.to_excel(path, index=False)


def convert_list_to_string_with_comma(product_data: dict) -> dict:
    output_dict = {}

    for key, value in product_data.items():
        output_dict[key] = ", ".join(value)

    return output_dict


def convert_doc_to_docx(path_to_file_doc: Path) -> Path:
    path_to_file_docx = Path(path_to_file_doc.stem, ".docx")

    # Create an object of the Document class
    document = Document()
    # Load a Word DOC file
    document.LoadFromFile(str(path_to_file_doc))

    # Save the DOC file to DOCX format
    document.SaveToFile(str(path_to_file_docx), FileFormat.Docx2016)
    # Close the Document object
    document.Close()

    return path_to_file_docx